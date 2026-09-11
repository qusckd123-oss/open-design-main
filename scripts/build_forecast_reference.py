import argparse
import json
import re
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
CONFIG_PATH = ROOT / "config" / "forecast-v1.json"
SOURCE_DIR = ROOT / ".local-forecast-reference"
OUTPUT_PATH = ROOT / "data" / "forecast-reference-v1.json"


def text(value):
    return str(value or "").strip()


def number(value):
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def style_category(style_code):
    match = re.match(r"^WA\d{4}([A-Z]{2})", text(style_code))
    return match.group(1) if match else ""


def gender_group(name):
    normalized = text(name).upper()
    return "WOMENS" if "우먼스" in normalized or "WOMEN" in normalized else "UNISEX"


def iso_date(value):
    if isinstance(value, (date, datetime)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return text(value)[:10]


def load_rows(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    header_dates = [iso_date(cell.value) if cell.value is not None else "" for cell in sheet[1]]
    headers = [text(cell.value) for cell in sheet[2]]
    rows = list(sheet.iter_rows(min_row=3, values_only=True))
    workbook.close()
    return header_dates, headers, rows


def build_reference(config):
    weekly_path = SOURCE_DIR / config["sourceFiles"]["weeklySales"]
    inbound_path = SOURCE_DIR / config["sourceFiles"]["inbound"]
    for path in (weekly_path, inbound_path):
        if not path.exists():
            raise FileNotFoundError(f"forecast reference source missing: {path}")

    date_headers, headers, source_rows = load_rows(weekly_path)
    _, inbound_headers, inbound_rows = load_rows(inbound_path)
    app_categories = set(config["appCategories"])
    clean = config["cleanAnalog"]
    december_date = clean["decemberProxyDate"]
    week_columns = [(index, value) for index, value in enumerate(date_headers) if re.match(r"^\d{4}-\d{2}-\d{2}$", value)]
    december_column = next((index for index, value in week_columns if value == december_date), None)
    if december_column is None:
        raise ValueError(f"December proxy column not found: {december_date}")

    app_rows = [row for row in source_rows if style_category(row[0]) in app_categories]
    domestic_adult = [
        row for row in app_rows
        if not any(marker.upper() in text(row[1]).upper() for marker in clean["specialMarketMarkers"])
        and not any(marker.upper() in text(row[1]).upper() for marker in clean["nonAdultMarkers"])
    ]

    inbound_index = {name: index for index, name in enumerate(inbound_headers)}
    reorder_skus = {
        (text(row[inbound_index["품번"]]), text(row[inbound_index["색상"]]))
        for row in inbound_rows
        if number(row[inbound_index["차수"]]) >= 2
    }

    filtered = []
    exclusion_counts = Counter()
    for row in domestic_adult:
        if number(row[4]) != number(row[5]):
            exclusion_counts["priceChanged"] += 1
            continue
        if (text(row[0]), text(row[2])) in reorder_skus:
            exclusion_counts["explicitReorder"] += 1
            continue
        first_period_column = week_columns[0][0]
        if number(row[first_period_column + 1]) > number(row[first_period_column]):
            exclusion_counts["salesBeforeAnalysis"] += 1
            continue
        if number(row[december_column + 1]) < number(clean["minimumDecemberCumulativeSales"]):
            exclusion_counts["lowDecemberSales"] += 1
            continue
        filtered.append(row)

    supply_constrained_styles = {
        text(row[0]) for row in filtered
        if number(row[8]) > 0 and number(row[9]) / number(row[8]) < number(clean["minimumReceiptToOrderRatio"])
    }
    clean_rows = [row for row in filtered if text(row[0]) not in supply_constrained_styles]
    exclusion_counts["suspectedSupplyConstraint"] = len(filtered) - len(clean_rows)

    by_style = defaultdict(list)
    for row in clean_rows:
        by_style[text(row[0])].append(row)

    warnings = []
    styles = []
    for style_code, rows in sorted(by_style.items()):
        final_qty = sum(number(row[december_column + 1]) for row in rows)
        cumulative = []
        first_positive_index = None
        previous = 0.0
        corrected = False
        for sequence_index, (column, period_date) in enumerate(week_columns):
            if period_date > december_date:
                break
            period_qty = sum(number(row[column]) for row in rows)
            if first_positive_index is None and period_qty > 0:
                first_positive_index = sequence_index
            raw_cumulative = sum(number(row[column + 1]) for row in rows)
            share = raw_cumulative / final_qty if final_qty > 0 else 0.0
            share = max(0.0, min(1.0, share))
            if share < previous:
                share = previous
                corrected = True
            previous = share
            if first_positive_index is not None and sequence_index >= first_positive_index:
                cumulative.append(round(share, 6))
        if corrected:
            warnings.append(f"monotonic correction applied: {style_code}")
        if not cumulative or any(cumulative[index] < cumulative[index - 1] for index in range(1, len(cumulative))):
            raise ValueError(f"non-monotonic style curve: {style_code}")
        first_date = week_columns[first_positive_index][1]
        first_date_value = datetime.fromisoformat(first_date)
        styles.append({
            "styleCode": style_code,
            "name": text(rows[0][1]),
            "category": style_category(style_code),
            "genderGroup": gender_group(rows[0][1]),
            "price": round(number(rows[0][5]), 2),
            "firstPositiveSalesDate": first_date,
            "launchTimingIndex": first_date_value.month * 5 + ((first_date_value.day - 1) // 7 + 1),
            "colorSkuCount": len(rows),
            "decemberCumulativeSales": round(final_qty),
            "cumulativeShareBySellingWeek": cumulative,
        })

    category_styles = defaultdict(list)
    for row in styles:
        category_styles[row["category"]].append(row)
    category_curves = {}
    for category, category_rows in sorted(category_styles.items()):
        max_weeks = max(len(row["cumulativeShareBySellingWeek"]) for row in category_rows)
        curve = []
        previous = 0.0
        for week_index in range(max_weeks):
            values = [row["cumulativeShareBySellingWeek"][week_index] for row in category_rows if week_index < len(row["cumulativeShareBySellingWeek"])]
            raw = sum(values) / len(values) if values else previous
            corrected = max(previous, raw)
            if corrected > raw:
                warnings.append(f"category monotonic correction applied: {category} W{week_index + 1}")
            curve.append(round(corrected, 6))
            previous = corrected
        if any(curve[index] < curve[index - 1] for index in range(1, len(curve))):
            raise ValueError(f"non-monotonic category curve: {category}")
        category_curves[category] = curve

    category_sku_sample_sizes = Counter(style_category(row[0]) for row in clean_rows)

    counts = {
        "app": {"sku": len(app_rows), "style": len({text(row[0]) for row in app_rows})},
        "domesticAdultApp": {"sku": len(domestic_adult), "style": len({text(row[0]) for row in domestic_adult})},
        "cleanAnalog": {"sku": len(clean_rows), "style": len(styles)},
    }
    expected = {"app": {"sku": 487, "style": 222}, "domesticAdultApp": {"sku": 462, "style": 207}, "cleanAnalog": {"sku": 348, "style": 166}}
    if counts != expected:
        raise ValueError(f"reference population mismatch: expected={expected}, actual={counts}")

    return {
        "meta": {
            "referenceVersion": config["referenceVersion"],
            "generatedAt": datetime.now().astimezone().isoformat(timespec="seconds"),
            "sourceScope": "25FW APP",
            "decemberProxyDate": december_date,
            "counts": counts,
            "exclusionCounts": dict(exclusion_counts),
            "supplyConstrainedStyles": sorted(supply_constrained_styles),
            "warnings": warnings,
        },
        "categorySampleSizes": dict(sorted(category_sku_sample_sizes.items())),
        "categoryStyleSampleSizes": {category: len(rows) for category, rows in sorted(category_styles.items())},
        "categoryCumulativeCurves": category_curves,
        "styles": styles,
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--validate-only", action="store_true")
    args = parser.parse_args()
    config = json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
    payload = build_reference(config)
    if not args.validate_only:
        OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
        OUTPUT_PATH.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"ok": True, "output": str(OUTPUT_PATH), "counts": payload["meta"]["counts"], "warnings": len(payload["meta"]["warnings"]), "validateOnly": args.validate_only}, ensure_ascii=False))


if __name__ == "__main__":
    main()
