from __future__ import annotations

import json
import re
import argparse
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
WORKSPACE = ROOT.parents[1]
SALES_ROOT = WORKSPACE / "추가 데이터" / "판매 데이터"
DEFAULT_CURRENT_PERIOD = "260810~260816"
DEFAULT_PRIOR_PERIOD = "260803~260809"

CATEGORY_LABELS = {
    "ST": "반팔티셔츠",
    "SO": "숏팬츠",
    "CD": "카디건",
    "CA": "모자",
    "BP": "가방",
    "SS": "반팔셔츠",
    "PT": "팬츠",
    "HZ": "하프집업",
    "KT": "니트",
    "CR": "크루넥",
}


def load_product_images() -> dict[str, dict[str, str]]:
    image_file = ROOT / "data" / "product_images.json"
    if not image_file.exists():
        return {}
    return json.loads(image_file.read_text(encoding="utf-8"))


def style_category(style: str) -> str:
    match = re.match(r"^WA\d{4}([A-Z]{2})", style)
    return match.group(1) if match else "ETC"


def style_season(style: str) -> str:
    match = re.match(r"^WA(\d{2})(\d{2})", style)
    if not match:
        return "미분류"
    year, drop = match.groups()
    season = "FW" if int(drop) >= 3 else "SS"
    return f"{year}{season}"


def style_gender(name: str) -> str:
    return "WOMEN" if "우먼" in name or "우먼스" in name else "UNISEX"


def xlsx_in(folder: Path, starts_with: str) -> Path:
    matches = sorted(folder.glob(f"{starts_with}*.xlsx"))
    if not matches:
        raise FileNotFoundError(f"{folder} 안에서 {starts_with}*.xlsx 파일을 찾지 못했습니다.")
    return matches[0]


def read_week_sales(path: Path) -> tuple[Counter[str], dict[str, str], Counter[str]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    style_amount: Counter[str] = Counter()
    style_names: dict[str, str] = {}
    style_qty: Counter[str] = Counter()

    for row in sheet.iter_rows(min_row=3, values_only=True):
        style = row[4]
        if not isinstance(style, str) or not style.startswith("WA26"):
            continue

        amount = row[11] or 0
        qty = row[10] or 0
        style_amount[style] += float(amount)
        style_qty[style] += int(qty)
        style_names[style] = str(row[5] or style)

    return style_amount, style_names, style_qty


def read_inventory_for(current_dir: Path, styles: set[str]) -> dict[str, dict[str, Any]]:
    trend_file = xlsx_in(current_dir, "와키윌리_26SS 전상품 판매추이")
    workbook = openpyxl.load_workbook(trend_file, read_only=True, data_only=True)
    skip_sheets = {
        "TTL",
        "ERP RAW_260816",
        "26SS 기획현황",
        "상품기획안_종합",
        "회의양식(주간별 판매비교)",
        "회의양식(베스트10비교)",
        "24' 품번별판매추이_누계",
        "24' 품번별판매추이_기간",
    }
    found: dict[str, dict[str, Any]] = {}

    for sheet in workbook.worksheets:
        if sheet.title in skip_sheets or "판매추이" in sheet.title:
            continue
        for row in sheet.iter_rows(min_row=1, max_col=32, values_only=True):
            style = row[7] if len(row) > 7 else None
            if style not in styles:
                continue

            in_qty = float(row[19] or 0)
            cum_qty = float(row[20] or 0)
            sell_rate = float(row[23] or 0)
            found[style] = {
                "category": row[5] or style_category(style),
                "name": row[8] or style,
                "inQty": in_qty,
                "cumQty": cum_qty,
                "stock": max(0, int(round(in_qty - cum_qty))),
                "sellThrough": round(sell_rate * 100, 1),
            }
            if styles <= found.keys():
                return found

    return found


def pct_change(current: float, prior: float) -> float:
    if prior == 0:
        return 0.0
    return (current - prior) / prior * 100


def reorder_timing(sell_through: float) -> tuple[str, int]:
    if sell_through >= 60:
        return "30% 초과·긴급", 0
    if sell_through >= 30:
        return "30% 도달", 0
    if sell_through >= 25:
        return "30% 임박", round(30 - sell_through)
    return "30% 전 관찰", round(30 - sell_through)


def make_action(sales_m: float, wow: float | None, sell_through: float, stock: int) -> tuple[str, str, str]:
    stock_rate = max(0.0, 100 - sell_through)
    timing, gap = reorder_timing(sell_through)
    if sell_through >= 30 and sales_m >= 5:
        return "리오더 검토", "P1", f"{timing} 구간입니다. 리오더 투입 여부와 예상 입고 시점을 우선 확인"
    if sell_through >= 25 and sales_m >= 8:
        return "리오더 검토", "P2", f"{timing} 구간으로 30%까지 약 {gap}%p 남았습니다. 선제 리오더 검토"
    if wow is not None and wow <= -35 and stock_rate >= 65:
        return "프로모션 검토", "P3", "전주 대비 둔화와 높은 잔여재고율이 동시에 발생해 가격 할인/행사 검토"
    if stock >= 800 and stock_rate >= 55:
        return "배분/RT 검토", "P2", "잔여재고율과 절대 재고가 높아 매장 이동(RT) 또는 채널 추가 배분 검토"
    return "배분/RT 검토", "P3", "금주 판매 흐름과 매장별 재고 편차 기준으로 배분 유지"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Update latest.json from Wacky Willy weekly sales workbooks.")
    parser.add_argument("--current-period", default=DEFAULT_CURRENT_PERIOD, help="Current sales folder/period, e.g. 260810~260816.")
    parser.add_argument("--prior-period", default=DEFAULT_PRIOR_PERIOD, help="Prior comparison folder/period, e.g. 260803~260809.")
    parser.add_argument("--week-label", default="8월 2주차", help="Display label for the current period.")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    current_period = args.current_period
    prior_period = args.prior_period
    current_dir = SALES_ROOT / current_period
    prior_dir = SALES_ROOT / prior_period

    if not current_dir.exists():
        raise FileNotFoundError(f"Current period folder not found: {current_dir}")
    if not prior_dir.exists():
        raise FileNotFoundError(f"Prior period folder not found: {prior_dir}")

    current_sales, names, current_qty = read_week_sales(xlsx_in(current_dir, f"판매집계현황 {current_period}"))
    prior_sales, _, _ = read_week_sales(xlsx_in(prior_dir, f"판매집계현황 {prior_period}"))
    inventory = read_inventory_for(current_dir, set(current_sales.keys()))
    product_images = load_product_images()

    current_by_category: Counter[str] = Counter()
    prior_by_category: Counter[str] = Counter()
    stock_by_category: Counter[str] = Counter()
    sold_by_category: Counter[str] = Counter()

    for style, amount in current_sales.items():
        current_by_category[style_category(style)] += amount

    for style, amount in prior_sales.items():
        prior_by_category[style_category(style)] += amount

    for style, row in inventory.items():
        category = style_category(style)
        stock_by_category[category] += row["stock"]
        sold_by_category[category] += row["cumQty"]

    categories = []
    for category, _ in current_by_category.most_common(8):
        current_m = round(current_by_category[category] / 1_000_000, 1)
        prior_m = round(prior_by_category[category] / 1_000_000, 1)
        stock = int(stock_by_category[category])
        sold = float(sold_by_category[category])
        sell_through = round(sold / (sold + stock) * 100, 1) if sold + stock > 0 else 0
        categories.append(
            {
                "category": category,
                "categoryName": CATEGORY_LABELS.get(category, category),
                "season": "26SS/26FW",
                "sales": current_m,
                "target": prior_m,
                "sellThrough": sell_through,
                "stock": stock,
                "wow": round(pct_change(current_by_category[category], prior_by_category[category]), 1),
            }
        )

    styles = []
    for style, amount in current_sales.most_common(18):
        inv = inventory.get(style, {})
        sales_m = round(amount / 1_000_000, 1)
        prior = prior_sales.get(style, 0)
        wow = pct_change(amount, prior) if prior else None
        sell_through = float(inv.get("sellThrough", 0))
        stock = int(inv.get("stock", 0))
        stock_rate = round(max(0.0, 100 - sell_through), 1)
        timing_label, timing_gap = reorder_timing(sell_through)
        display_name = str(inv.get("name") or names.get(style) or style)
        action, priority, note = make_action(sales_m, wow, sell_through, stock)
        if wow is not None:
            note = f"{note} (전주 대비 {wow:+.1f}%)"

        image_info = product_images.get(style, {})

        styles.append(
            {
                "sku": style,
                "name": display_name,
                "category": style_category(style),
                "categoryName": CATEGORY_LABELS.get(style_category(style), style_category(style)),
                "season": style_season(style),
                "gender": style_gender(display_name),
                "sales": sales_m,
                "priorSales": round(prior / 1_000_000, 1),
                "quantity": int(current_qty[style]),
                "inQty": int(round(float(inv.get("inQty", 0)))),
                "cumQty": int(round(float(inv.get("cumQty", 0)))),
                "stock": stock,
                "sellThrough": sell_through,
                "stockRate": stock_rate,
                "reorderTiming": timing_label,
                "reorderGapTo30": timing_gap,
                "wow": round(wow, 1) if wow is not None else None,
                "action": action,
                "priority": priority,
                "note": note,
                "imageUrl": image_info.get("imageUrl"),
                "productUrl": image_info.get("productUrl"),
                "imageSource": image_info.get("source"),
                "imageSourceName": image_info.get("sourceName"),
            }
        )

    current_total = sum(current_sales.values())
    prior_total = sum(prior_sales.values())
    total_wow = pct_change(current_total, prior_total)
    best_category = max(categories, key=lambda row: row["sales"]) if categories else None
    reorder_styles = [row for row in styles if row["action"] == "리오더"]
    best_reorder = reorder_styles[0] if reorder_styles else (styles[0] if styles else None)

    payload = {
        "meta": {
            "brand": "Wacky Willy",
            "season": "26SS/26FW",
            "weekLabel": args.week_label,
            "period": current_period,
            "comparePeriod": prior_period,
            "amountUnit": "VAT- / 백만원",
            "updatedAt": datetime.now().strftime("%Y-%m-%d %H:%M"),
            "source": "판매집계현황, 전상품 판매추이",
            "targetLabel": "전주 매출",
        },
        "categories": categories,
        "styles": styles,
        "launch": [
            {"label": "전주", "value": round(prior_total / 1_000_000, 1)},
            {"label": "금주", "value": round(current_total / 1_000_000, 1)},
        ],
        "summary": {
            "headline": f"{args.week_label} WA26 주간 매출은 {current_total / 100_000_000:.2f}억으로 전주 대비 {total_wow:+.1f}%입니다.",
            "message": (
                f"{best_category['category']}({best_category['categoryName']})가 금주 매출을 가장 크게 견인했습니다. "
                f"상위 스타일 중 {best_reorder['sku']}는 판매율 30% 리오더 시점과 잔여재고율 기준으로 우선 점검하고, "
                "전주 대비 하락한 고재고율 스타일은 배분/RT 또는 프로모션 검토를 분리해서 보겠습니다."
            )
            if best_category and best_reorder
            else "금주 판매 데이터가 충분하지 않습니다.",
        },
    }

    output_paths = [
        ROOT / "data" / "latest.json",
        ROOT / "public" / "data" / "latest.json",
    ]
    for output_path in output_paths:
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    archive_name = f"{current_period.replace('~', '-')}.json"
    for archive_dir in [ROOT / "data" / "archive", ROOT / "public" / "data" / "archive"]:
        archive_dir.mkdir(parents=True, exist_ok=True)
        (archive_dir / archive_name).write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    print(json.dumps({"currentTotal": current_total, "priorTotal": prior_total, "styles": len(styles), "categories": len(categories)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
