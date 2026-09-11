#!/usr/bin/env python3
"""Audit 26FW order rows for overseas PO applicability.

This module is deliberately read-only with respect to the production SKU fact.
It proves which PO-row order quantities can be removed from a domestic order
fact and reports, but does not apply, changes that would alter Analog Pace.
"""

from __future__ import annotations

import argparse
import json
import math
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SKU = ROOT / "data" / "sku-latest.json"
DEFAULT_ANALOG = ROOT / "data" / "sku-analog-pace-coverage-diagnostic.json"
DEFAULT_OUTPUT = ROOT / "data" / "sku-overseas-po-applicability.json"

KEYWORD_CLASSIFICATION = (
    ("대만", "TAIWAN"),
    ("일본", "JAPAN"),
    ("글로벌", "GLOBAL"),
    ("수주", "ORDER_ONLY"),
)
REQUIRED_HEADERS = ("발주번호", "품번", "색상", "발주수량")


def text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def number(value: Any) -> float:
    try:
        parsed = float(value or 0)
    except (TypeError, ValueError):
        return 0.0
    return parsed if math.isfinite(parsed) else 0.0


def canonical(value: Any) -> str:
    return text(value).upper().replace(" ", "").replace("-", "")


def classify_po_number(value: Any) -> dict[str, Any]:
    po_number = text(value)
    matched = [keyword for keyword, _ in KEYWORD_CLASSIFICATION if keyword in po_number]
    classification = next(
        (label for keyword, label in KEYWORD_CLASSIFICATION if keyword in matched),
        "DOMESTIC_OR_OTHER",
    )
    return {
        "poNumber": po_number,
        "isOverseasPo": bool(matched),
        "classification": classification,
        "matchedKeywords": matched,
    }


def classify_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    classified = []
    for index, source in enumerate(rows, start=1):
        style = canonical(source.get("styleCode"))
        color = canonical(source.get("colorCode"))
        classified.append(
            {
                **source,
                "sourceRow": source.get("sourceRow", index),
                "styleCode": style,
                "colorCode": color,
                "sku": f"{style}{color}",
                "orderQty": number(source.get("orderQty")),
                **classify_po_number(source.get("poNumber")),
            }
        )
    return classified


def find_header(ws: Any) -> tuple[int, dict[str, int]]:
    for row_number in range(1, min(ws.max_row, 30) + 1):
        columns = {
            text(ws.cell(row_number, column).value): column
            for column in range(1, ws.max_column + 1)
            if text(ws.cell(row_number, column).value)
        }
        if all(header in columns for header in REQUIRED_HEADERS):
            return row_number, columns
    raise ValueError(f"required headers not found: {', '.join(REQUIRED_HEADERS)}")


def parse_workbook(path: Path) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    workbook = load_workbook(path, data_only=False, read_only=False)
    sheets = workbook.sheetnames
    if len(workbook.worksheets) != 1:
        raise ValueError(f"expected one order sheet, found {len(workbook.worksheets)}")
    ws = workbook.worksheets[0]
    header_row, columns = find_header(ws)
    inbound_header = next(
        (name for name in ("입고수량", "누적입고수량", "실입고수량") if name in columns),
        None,
    )
    rows: list[dict[str, Any]] = []
    blank_rows: list[int] = []
    repeated_headers: list[int] = []
    for row_number in range(header_row + 1, ws.max_row + 1):
        values = [ws.cell(row_number, column).value for column in range(1, ws.max_column + 1)]
        if not any(value not in (None, "") for value in values):
            blank_rows.append(row_number)
            continue
        if text(ws.cell(row_number, columns["품번"]).value) == "품번":
            repeated_headers.append(row_number)
            continue
        rows.append(
            {
                "sourceRow": row_number,
                "styleCode": ws.cell(row_number, columns["품번"]).value,
                "colorCode": ws.cell(row_number, columns["색상"]).value,
                "poNumber": ws.cell(row_number, columns["발주번호"]).value,
                "orderQty": ws.cell(row_number, columns["발주수량"]).value,
                "inboundQty": (
                    ws.cell(row_number, columns[inbound_header]).value
                    if inbound_header
                    else None
                ),
            }
        )
    meta = {
        "sourceFile": path.name,
        "sheetNames": sheets,
        "sheetName": ws.title,
        "maxRow": ws.max_row,
        "maxColumn": ws.max_column,
        "headerRow": header_row,
        "columns": {
            "styleCode": {"name": "품번", "index": columns["품번"]},
            "colorCode": {"name": "색상", "index": columns["색상"]},
            "poNumber": {"name": "발주번호", "index": columns["발주번호"]},
            "orderQty": {"name": "발주수량", "index": columns["발주수량"]},
            "inboundQty": {"name": inbound_header, "index": columns.get(inbound_header)},
        },
        "relatedInboundFields": [name for name in ("납기완료일", "입고마감") if name in columns],
        "mergedRanges": [str(cell_range) for cell_range in ws.merged_cells.ranges],
        "blankRows": blank_rows,
        "repeatedHeaderRows": repeated_headers,
    }
    return classify_rows(rows), meta


def flatten_app_skus(payload: dict[str, Any]) -> dict[str, dict[str, Any]]:
    return {
        row["sku"]: row
        for style in payload.get("styles", {}).values()
        for row in style.get("skus", [])
        if row.get("season") == "26FW" and row.get("productGroup") == "APP"
    }


def build_report(
    rows: list[dict[str, Any]],
    sku_payload: dict[str, Any],
    analog_payload: dict[str, Any],
    workbook: dict[str, Any] | None = None,
) -> dict[str, Any]:
    classified = rows if all("isOverseasPo" in row for row in rows) else classify_rows(rows)
    overseas = [row for row in classified if row["isOverseasPo"]]
    by_sku: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in classified:
        if row["sku"]:
            by_sku[row["sku"]].append(row)

    app = flatten_app_skus(sku_payload)
    affected = sorted(set(row["sku"] for row in overseas) & set(app))
    aggregates: dict[str, dict[str, Any]] = {}
    mixed_skus: list[str] = []
    overseas_only_skus: list[str] = []
    for sku, sku_rows in by_sku.items():
        total = sum(row["orderQty"] for row in sku_rows)
        overseas_qty = sum(row["orderQty"] for row in sku_rows if row["isOverseasPo"])
        domestic_qty = total - overseas_qty
        aggregates[sku] = {
            "excelTotalOrderQty": total,
            "domesticOrderQty": domestic_qty,
            "overseasOrderQty": overseas_qty,
        }
        if overseas_qty:
            if any(not row["isOverseasPo"] for row in sku_rows):
                mixed_skus.append(sku)
            else:
                overseas_only_skus.append(sku)

    reconciliation = []
    for sku in sorted(set(app) & set(aggregates)):
        aggregate = aggregates[sku]
        current = number(app[sku].get("orderQty"))
        reconciliation.append(
            {
                "sku": sku,
                **aggregate,
                "currentOrderQty": current,
                "currentIncludesOverseasOrder": abs(current - aggregate["excelTotalOrderQty"]) < 1e-9,
            }
        )

    analog_rows = {
        row["sku"]: row
        for row in analog_payload.get("current", {}).get("rows", [])
    }
    analog_impact = []
    for sku in affected:
        current = analog_rows.get(sku)
        if not current or current.get("primaryFailureReason") != "PACE_READY":
            continue
        aggregate = aggregates[sku]
        old_index = current.get("analogPaceIndex")
        domestic = aggregate["domesticOrderQty"]
        expected_index = (
            number(old_index) * aggregate["excelTotalOrderQty"] / domestic
            if old_index is not None and domestic > 0
            else None
        )
        analog_impact.append(
            {
                "sku": sku,
                "styleCode": current.get("styleCode"),
                "sellingWeek": current.get("sellingWeek"),
                "currentOrderQty": aggregate["excelTotalOrderQty"],
                "domesticOrderQty": domestic,
                "overseasOrderQty": aggregate["overseasOrderQty"],
                "currentPaceIndex": old_index,
                "expectedPaceIndexIfOrderWereReplaced": expected_index,
                "currentPacePercentile": current.get("analogPacePercentile"),
            }
        )

    classification_rows = Counter(row["classification"] for row in overseas)
    classification_qty = Counter()
    matched_rows = Counter()
    matched_qty = Counter()
    po_examples = Counter()
    for row in overseas:
        classification_qty[row["classification"]] += row["orderQty"]
        po_examples[row["poNumber"]] += 1
        for keyword in row["matchedKeywords"]:
            matched_rows[keyword] += 1
            matched_qty[keyword] += row["orderQty"]

    special_skus = (
        "WA2603CRT1BK",
        "WA2603CRT1GR",
        "WA2603STT1BK",
        "WA2603STT1WH",
        "WA2603STT2CH",
    )
    special = []
    for sku in special_skus:
        sku_rows = by_sku.get(sku, [])
        matched = sorted({keyword for row in sku_rows for keyword in row["matchedKeywords"]})
        special.append(
            {
                "sku": sku,
                "isSpecialMarket": app.get(sku, {}).get("isSpecialMarket"),
                "overseasPoStatus": "OVERSEAS_PO_CONFIRMED" if matched else "UNKNOWN",
                "matchedKeywords": matched,
                "sourceRows": [row["sourceRow"] for row in sku_rows if row["isOverseasPo"]],
                "poNumbers": sorted({row["poNumber"] for row in sku_rows if row["isOverseasPo"]}),
                **aggregates.get(sku, {"excelTotalOrderQty": 0, "domesticOrderQty": 0, "overseasOrderQty": 0}),
            }
        )

    return {
        "schemaVersion": "sku-overseas-po-applicability-v1",
        "readOnlyDiagnostic": True,
        "productionBehaviorChanged": False,
        "source": workbook or {},
        "classificationContract": {
            "unit": "PO_ROW",
            "keywords": [keyword for keyword, _ in KEYWORD_CLASSIFICATION],
            "classificationPrecedence": [label for _, label in KEYWORD_CLASSIFICATION],
            "matchedKeywordsCanOverlap": True,
            "rowCountIsDeduplicated": True,
            "noSkuOrStylePropagation": True,
            "isSpecialMarketIsNotRouteEvidence": True,
        },
        "applicability": {
            "order": "CONFIRMED_EXCLUDABLE_FROM_DOMESTIC_ORDER",
            "inboundRule": "CONFIRMED_EXCLUDABLE_FROM_DOMESTIC_INBOUND",
            "inbound": "NEEDS_ADDITIONAL_SOURCE",
            "sales": "DOMESTIC_SALES_EXCLUSION_NOT_PROVEN",
            "erpStock": "NEEDS_ADDITIONAL_SOURCE",
            "reason": "The order workbook has PO-row order quantity but no PO-row inbound, sales, or ERP stock quantity.",
        },
        "summary": {
            "dataRowCount": len(classified),
            "overseasPoRowCount": len(overseas),
            "overseasPoSkuCount": len({row["sku"] for row in overseas}),
            "overseasPoStyleCount": len({row["styleCode"] for row in overseas}),
            "exclusiveClassificationRowCounts": {
                label: classification_rows.get(label, 0)
                for _, label in KEYWORD_CLASSIFICATION
            },
            "exclusiveClassificationOrderQty": {
                label: classification_qty.get(label, 0)
                for _, label in KEYWORD_CLASSIFICATION
            },
            "matchedKeywordRowCounts": {
                keyword: matched_rows.get(keyword, 0)
                for keyword, _ in KEYWORD_CLASSIFICATION
            },
            "matchedKeywordOrderQty": {
                keyword: matched_qty.get(keyword, 0)
                for keyword, _ in KEYWORD_CLASSIFICATION
            },
            "multiKeywordRowCount": sum(len(row["matchedKeywords"]) > 1 for row in overseas),
            "overseasPoOrderQty": sum(row["orderQty"] for row in overseas),
            "mixedDomesticAndOverseasSkuCount": len(mixed_skus),
            "overseasOnlySkuCount": len(overseas_only_skus),
            "app26FwSkuCount": len(app),
            "app26FwPresentInWorkbook": len(set(app) & set(aggregates)),
            "app26FwMissingFromWorkbook": sorted(set(app) - set(aggregates)),
            "app26FwAffectedSkuCount": len(affected),
            "app26FwAffectedStyleCount": len({app[sku]["styleCode"] for sku in affected}),
            "affectedCurrentOrderExactTotalMatches": sum(
                row["sku"] in affected and row["currentIncludesOverseasOrder"]
                for row in reconciliation
            ),
        },
        "poNumberExamples": dict(sorted(po_examples.items())),
        "overseasPoRows": [
            {
                key: row.get(key)
                for key in (
                    "sourceRow",
                    "styleCode",
                    "colorCode",
                    "sku",
                    "poNumber",
                    "orderQty",
                    "classification",
                    "matchedKeywords",
                )
            }
            for row in overseas
        ],
        "mixedDomesticAndOverseasSkus": sorted(mixed_skus),
        "overseasOnlySkus": sorted(overseas_only_skus),
        "app26FwAffectedSkus": affected,
        "currentOrderReconciliation": reconciliation,
        "specialMarketFive": special,
        "hypotheticalImpact": {
            "appSkuCountBefore": len(app),
            "appSkuCountAfter": len(app),
            "orderQtyChangedSkuCount": len(affected),
            "orderQtyReduction": sum(aggregates[sku]["overseasOrderQty"] for sku in affected),
            "inboundQtyChangedSkuCount": 0,
            "erpStockQtyChangedSkuCount": 0,
            "stockCoverChangedSkuCount": 0,
            "lifecycleChangedSkuCount": 0,
            "salesVelocityChangedSkuCount": 0,
            "analogPaceReadyAffectedSkuCount": len(analog_impact),
            "analogPaceIndexWouldChangeSkuCount": sum(
                row["domesticOrderQty"] > 0 for row in analog_impact
            ),
            "analogPaceWouldLoseOrderDenominatorSkuCount": sum(
                row["domesticOrderQty"] <= 0 for row in analog_impact
            ),
            "analogPaceAffectedRows": analog_impact,
            "implementationStatus": "STOPPED_BEFORE_PRODUCTION_ORDER_REPLACEMENT",
            "stopReason": "Replacing orderQty would change protected ORDER-denominator Analog Pace; methodology approval is required.",
        },
    }


def load_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--orders", type=Path, required=True)
    parser.add_argument("--sku", type=Path, default=DEFAULT_SKU)
    parser.add_argument("--analog", type=Path, default=DEFAULT_ANALOG)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    rows, workbook = parse_workbook(args.orders)
    report = build_report(rows, load_json(args.sku), load_json(args.analog), workbook)
    args.output.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report["summary"], ensure_ascii=False))


if __name__ == "__main__":
    main()
